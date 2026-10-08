#!/usr/bin/env python3
"""PPT 几何 QA: 检查所有 shape 是否越出 16:9 画布(10x5.625in) + dump 每页文本。
用法: python3 qa_geom.py <deck1.pptx> [deck2.pptx ...]
依赖: pip3 install --break-system-packages --user python-pptx
"""
import sys
from pptx import Presentation
from pptx.util import Emu

SLIDE_W = 10.0   # inches (16:9)
SLIDE_H = 5.625

def inch(v): return Emu(v).inches

def check(path):
    prs = Presentation(path)
    issues = []
    print(f"\n===== {path} =====")
    print(f"slides: {len(prs.slides)}")
    for si, slide in enumerate(prs.slides, 1):
        shapes = list(slide.shapes)
        texts = []
        for sh in shapes:
            try:
                l, t, w, h = inch(sh.left), inch(sh.top), inch(sh.width), inch(sh.height)
            except Exception:
                continue
            if l < -0.02 or t < -0.02 or l + w > SLIDE_W + 0.02 or t + h > SLIDE_H + 0.02:
                issues.append(f"S{si}: 越界 shape '{sh.shape_type}' at ({l:.2f},{t:.2f}) size ({w:.2f}x{h:.2f}) -> right={l+w:.2f} bottom={t+h:.2f}")
            if sh.has_text_frame and sh.text_frame.text.strip():
                texts.append(sh.text_frame.text.strip().replace('\n', ' / '))
        print(f"--- Slide {si} ({len(shapes)} shapes) ---")
        for t in texts[:5]:
            print("   ", t[:70])
    return issues

all_issues = []
for p in sys.argv[1:]:
    all_issues += check(p)

print("\n\n===== 越界问题汇总 =====")
if all_issues:
    for i in all_issues: print("❌", i)
    sys.exit(1)
print("✅ 无越界问题")
